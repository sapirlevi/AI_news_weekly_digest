"""Build the weekly .pptx slide deck from digest.json and the raw sources.

Embeds matplotlib charts and YouTube thumbnails. Designed to be robust —
if any optional asset fetch fails, the slide still renders.
"""
import io
import sys
from collections import Counter
from datetime import datetime

import requests

from util_io import TMP, read_json, now_utc

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("ERROR: pip install python-pptx", file=sys.stderr)
    sys.exit(1)

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
except ImportError:
    print("ERROR: pip install matplotlib", file=sys.stderr)
    sys.exit(1)


NAVY = RGBColor(0x0B, 0x1F, 0x3A)
ACCENT = RGBColor(0xFF, 0x6B, 0x35)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
BG = RGBColor(0xF7, 0xF7, 0xF7)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def blank_slide(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    # background
    bg = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
    bg.line.fill.background()
    bg.fill.solid()
    bg.fill.fore_color.rgb = BG
    return s


def add_text(slide, text, left, top, width, height, *, size=18, bold=False,
             color=NAVY, align=None, hyperlink=None):
    from pptx.enum.text import PP_ALIGN
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    if align == "center":
        p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = "Calibri"
    if hyperlink:
        run.hyperlink.address = hyperlink
        run.font.underline = True
    return box


def add_bullets(slide, items, left, top, width, height, *, size=16, color=NAVY):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run = p.add_run()
        run.text = f"•  {it}"
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.name = "Calibri"
        p.space_after = Pt(6)


def add_header_bar(slide, title, subtitle=""):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(0.9))
    bar.line.fill.background()
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    add_text(slide, title, Inches(0.5), Inches(0.15), Inches(9), Inches(0.5),
             size=24, bold=True, color=WHITE)
    if subtitle:
        add_text(slide, subtitle, Inches(0.5), Inches(0.52), Inches(12), Inches(0.4),
                 size=12, color=RGBColor(0xCC, 0xD4, 0xE0))


def accent_stripe(slide):
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(0.9), SLIDE_W, Inches(0.08))
    stripe.line.fill.background()
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ACCENT


# ---------- slides ----------

def cover_slide(prs, week_start, week_end):
    s = blank_slide(prs)
    # hero block
    hero = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, Inches(4.2))
    hero.line.fill.background()
    hero.fill.solid()
    hero.fill.fore_color.rgb = NAVY
    stripe = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(4.2), SLIDE_W, Inches(0.12))
    stripe.line.fill.background()
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = ACCENT
    add_text(s, "AI INDUSTRY WEEKLY", Inches(0.8), Inches(1.2), Inches(12), Inches(0.8),
             size=48, bold=True, color=WHITE)
    add_text(s, f"Week of {week_start} – {week_end}",
             Inches(0.8), Inches(2.2), Inches(12), Inches(0.6),
             size=22, color=RGBColor(0xCC, 0xD4, 0xE0))
    add_text(s, "News · YouTube trends · Prompts · Tools",
             Inches(0.8), Inches(2.9), Inches(12), Inches(0.5),
             size=16, color=RGBColor(0xCC, 0xD4, 0xE0))
    add_text(s, "Automated briefing — WAT framework",
             Inches(0.8), Inches(5.5), Inches(12), Inches(0.4),
             size=12, color=MUTED)


def summary_slide(prs, bullets):
    s = blank_slide(prs)
    add_header_bar(s, "Executive Summary", "What you need to know this week")
    accent_stripe(s)
    add_bullets(s, bullets or ["Quiet week — no major developments."],
                Inches(0.8), Inches(1.4), Inches(12), Inches(5.5), size=20)


def top_stories_slides(prs, stories):
    if not stories:
        return
    # Split into groups of 4 stories per slide
    chunks = [stories[i:i + 4] for i in range(0, len(stories), 4)]
    for idx, chunk in enumerate(chunks, 1):
        s = blank_slide(prs)
        add_header_bar(s, f"Top Stories ({idx}/{len(chunks)})",
                       "Major announcements from AI providers")
        accent_stripe(s)
        top = Inches(1.25)
        for st in chunk:
            card_h = Inches(1.35)
            card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.5), top,
                                      Inches(12.3), card_h)
            card.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
            card.fill.solid()
            card.fill.fore_color.rgb = WHITE
            add_text(s, (st.get("provider") or "").upper(),
                     Inches(0.8), top + Inches(0.1),
                     Inches(4), Inches(0.3), size=11, bold=True, color=ACCENT)
            add_text(s, st.get("headline", ""),
                     Inches(0.8), top + Inches(0.35),
                     Inches(12), Inches(0.5), size=16, bold=True, color=NAVY,
                     hyperlink=st.get("url") or None)
            add_text(s, st.get("why_it_matters", ""),
                     Inches(0.8), top + Inches(0.78),
                     Inches(12), Inches(0.55), size=12, color=MUTED)
            top += card_h + Inches(0.1)


def trending_videos_slide(prs, vids, videos_by_id):
    if not vids:
        return
    s = blank_slide(prs)
    add_header_bar(s, "Trending on YouTube", "From your curated channel list")
    accent_stripe(s)
    # 3x3 grid of cards with thumbnails (9 videos)
    grid = vids[:9]
    col_w = Inches(4.15)
    row_h = Inches(1.9)
    gap = Inches(0.1)
    start_left = Inches(0.45)
    start_top = Inches(1.25)
    for i, v in enumerate(grid):
        r, c = divmod(i, 3)
        left = start_left + c * (col_w + gap)
        top = start_top + r * (row_h + gap)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, col_w, row_h)
        card.line.color.rgb = RGBColor(0xDD, 0xDD, 0xDD)
        card.fill.solid()
        card.fill.fore_color.rgb = WHITE

        # thumbnail
        thumb_url = (videos_by_id.get(v.get("videoId", ""), {})
                     .get("thumbnailUrl", ""))
        thumb_h = Inches(0.85)
        if thumb_url:
            try:
                img = requests.get(thumb_url, timeout=10).content
                thumb_w = Inches(0.85 * 16 / 9)  # preserve 16:9 aspect at thumb_h
                thumb_left = left + (col_w - thumb_w) / 2
                s.shapes.add_picture(io.BytesIO(img), thumb_left,
                                     top + Inches(0.08), height=thumb_h)
            except Exception:
                pass
        text_top = top + thumb_h + Inches(0.18)
        vid_id = v.get("videoId", "")
        yt_url = f"https://youtu.be/{vid_id}" if vid_id else None
        add_text(s, v.get("title", "")[:80],
                 left + Inches(0.12), text_top,
                 col_w - Inches(0.24), Inches(0.35), size=9, bold=True, color=NAVY,
                 hyperlink=yt_url)
        add_text(s, f"{v.get('channel','')}  ·  {v.get('views',0):,} views",
                 left + Inches(0.12), text_top + Inches(0.35),
                 col_w - Inches(0.24), Inches(0.2), size=8, color=MUTED)
        add_text(s, (v.get("key_takeaway", "") or "")[:100],
                 left + Inches(0.12), text_top + Inches(0.55),
                 col_w - Inches(0.24), Inches(0.3), size=8, color=NAVY)


def tips_slide(prs, title, subtitle, items, formatter):
    if not items:
        return
    s = blank_slide(prs)
    add_header_bar(s, title, subtitle)
    accent_stripe(s)
    bullets = [formatter(it) for it in items]
    add_bullets(s, bullets, Inches(0.7), Inches(1.3), Inches(12.2), Inches(5.8), size=14)


def tools_and_tips_slide(prs, items):
    if not items:
        return
    per_page = 6
    pages = [items[i:i + per_page] for i in range(0, len(items), per_page)]
    total = len(pages)
    for idx, chunk in enumerate(pages, 1):
        s = blank_slide(prs)
        suffix = f"  ({idx}/{total})" if total > 1 else ""
        add_header_bar(s,
                       f"Tools, Prompts & Setup Tips Worth Trying{suffix}",
                       "Hands-on recommendations — click any tip to jump to the source")
        accent_stripe(s)
        top = Inches(1.3)
        left = Inches(0.7)
        width = Inches(12.2)
        for it in chunk:
            url = it.get("url") or None
            title = it.get("title", "")
            what = it.get("what", "")
            source = it.get("source", "")
            add_text(s, f"•  {title}", left, top, width, Inches(0.34),
                     size=14, bold=True, color=NAVY, hyperlink=url)
            sub_parts = [p for p in [what, f"({source})" if source else ""] if p]
            sub = "  ".join(sub_parts)
            add_text(s, sub, left + Inches(0.25), top + Inches(0.32),
                     width - Inches(0.25), Inches(0.5), size=11, color=MUTED)
            top += Inches(0.95)


def qa_ai_slide(prs, items):
    if not items:
        return
    per_page = 6
    pages = [items[i:i + per_page] for i in range(0, len(items), per_page)]
    total = len(pages)
    for idx, chunk in enumerate(pages, 1):
        s = blank_slide(prs)
        suffix = f"  ({idx}/{total})" if total > 1 else ""
        add_header_bar(s,
                       f"QA × AI — Tools, Frameworks & How-Tos{suffix}",
                       "Actionable picks for QA leaders — click any title to read the post")
        accent_stripe(s)
        top = Inches(1.3)
        left = Inches(0.7)
        width = Inches(12.2)
        for it in chunk:
            url = it.get("url") or None
            title = it.get("title", "")
            what = it.get("what", "")
            source = it.get("source", "")
            add_text(s, f"•  {title}", left, top, width, Inches(0.34),
                     size=14, bold=True, color=NAVY, hyperlink=url)
            sub_parts = [p for p in [what, f"({source})" if source else ""] if p]
            sub = "  ".join(sub_parts)
            add_text(s, sub, left + Inches(0.25), top + Inches(0.32),
                     width - Inches(0.25), Inches(0.5), size=11, color=MUTED)
            top += Inches(0.95)


def chart_png(fig):
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf


def charts_slide(prs, videos, news_items):
    s = blank_slide(prs)
    add_header_bar(s, "Industry Signals", "Quantitative view of the week")
    accent_stripe(s)

    # Chart 1: top channels by views
    by_channel = Counter()
    for v in videos:
        by_channel[v.get("channel", "?")] += v.get("views", 0) or 0
    top_ch = by_channel.most_common(8)
    if top_ch:
        fig, ax = plt.subplots(figsize=(6.2, 3.6))
        labels = [c[:22] for c, _ in top_ch][::-1]
        values = [v for _, v in top_ch][::-1]
        ax.barh(labels, values, color="#0B1F3A")
        ax.set_title("Views per channel (this week)", fontsize=11, weight="bold")
        ax.tick_params(labelsize=9)
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        png = chart_png(fig)
        s.shapes.add_picture(png, Inches(0.3), Inches(1.3), width=Inches(6.4))

    # Chart 2: news items per provider
    by_prov = Counter(n.get("provider", "?") for n in news_items)
    if by_prov:
        fig, ax = plt.subplots(figsize=(6.2, 3.6))
        provs, counts = zip(*by_prov.most_common())
        ax.bar(provs, counts, color="#FF6B35")
        ax.set_title("News items per provider", fontsize=11, weight="bold")
        ax.tick_params(labelsize=9)
        for label in ax.get_xticklabels():
            label.set_rotation(30)
            label.set_ha("right")
        ax.spines["top"].set_visible(False)
        ax.spines["right"].set_visible(False)
        png = chart_png(fig)
        s.shapes.add_picture(png, Inches(6.8), Inches(1.3), width=Inches(6.2))


def sources_slide(prs, stories, vids, qa_items=None):
    s = blank_slide(prs)
    add_header_bar(s, "Sources", "Links for deeper reading")
    accent_stripe(s)
    entries = []
    for st in stories:
        if st.get("url"):
            label = f"[{st.get('provider','')}] {st.get('headline','')[:90]}"
            entries.append((label, st["url"]))
    for v in vids:
        if v.get("videoId"):
            label = f"[YT · {v.get('channel','')}] {v.get('title','')[:80]}"
            entries.append((label, f"https://youtu.be/{v['videoId']}"))
    for q in qa_items or []:
        if q.get("url"):
            label = f"[QA · {q.get('source','')}] {q.get('title','')[:80]}"
            entries.append((label, q["url"]))
    if not entries:
        add_text(s, "(no sources this week)", Inches(0.5), Inches(1.4),
                 Inches(12.4), Inches(0.4), size=10, color=MUTED)
        return
    line_h = Inches(0.32)
    top = Inches(1.25)
    for label, url in entries[:18]:
        add_text(s, label, Inches(0.5), top, Inches(12.4), line_h,
                 size=10, color=NAVY, hyperlink=url)
        top += line_h


def main():
    digest = read_json(TMP / "digest.json")
    videos_raw = read_json(TMP / "youtube_videos.json").get("videos", [])
    news_raw = read_json(TMP / "provider_news.json").get("items", [])
    videos_by_id = {v["videoId"]: v for v in videos_raw}

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    today = now_utc()
    week_start = (today.replace(hour=0, minute=0, second=0, microsecond=0)).strftime("%b %d")
    from datetime import timedelta
    week_start_dt = today - timedelta(days=7)
    week_start = week_start_dt.strftime("%b %d, %Y")
    week_end = today.strftime("%b %d, %Y")

    cover_slide(prs, week_start, week_end)
    summary_slide(prs, digest.get("executive_summary", []))
    top_stories_slides(prs, digest.get("top_stories", []))
    trending_videos_slide(prs, digest.get("trending_videos", []), videos_by_id)
    tools_and_tips_slide(prs, digest.get("tools_and_tips", []))
    qa_ai_slide(prs, digest.get("qa_ai", []))
    charts_slide(prs, videos_raw, news_raw)
    sources_slide(prs,
                  digest.get("top_stories", []),
                  digest.get("trending_videos", []),
                  digest.get("qa_ai", []))

    out_path = TMP / f"ai_digest_{today.strftime('%Y-%m-%d')}.pptx"
    prs.save(out_path)
    print(f"  wrote {out_path}")
    print(f"Done. {len(prs.slides)} slides.")


if __name__ == "__main__":
    main()
